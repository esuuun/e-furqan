import openpyxl, sys, os
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('Kamus Harf \'Amil.xlsx', data_only=True)
ws = wb['KAMUS KATA']
header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]
rows = list(ws.iter_rows(values_only=True))

words = {}
for r in rows[1:]:
    if not any(r): continue
    cat, no, word, arti, frek, surat, ayat = r[0], r[1], r[2], r[3], r[4], r[5], r[6]
    key = (str(cat), str(no))
    if key not in words:
        words[key] = {'cat': cat, 'no': no, 'word': word, 'arti': arti, 'frek': frek, 'verses': []}
    words[key]['verses'].append((surat, ayat))

print(f'Total Unique Words: {len(words)}')
for k, v in sorted(words.items(), key=lambda x: (x[0][0], int(x[0][1]) if str(x[0][1]).isdigit() else 99)):
    print(f"{v['cat']} | No {v['no']} | {v['word']} | {v['arti']} | Frek: {v['frek']} | {len(v['verses'])} ayat: {v['verses'][:3]}")
